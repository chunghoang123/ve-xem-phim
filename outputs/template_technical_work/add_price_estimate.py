from copy import copy
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

path = Path("../completed/Template-Technical-Completed.xlsx")
wb = load_workbook(path)
ws = wb["2. Price"]

ws["B5"] = "Dự án: AI Book Movie Tickets"
ws["B6"] = "Khách hàng: Cinex / Hệ thống rạp phim"
ws["E5"] = "Từ ngày : 15/07/2026"
ws["E6"] = "Đến ngày: 15/10/2026"

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
money_fill = PatternFill("solid", fgColor="E2F0D9")
subtotal_fill = PatternFill("solid", fgColor="D9EAD3")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

headers = {
    "G10": "Đơn giá/manday (VND)",
    "H10": "Kinh phí ước lượng (VND)",
    "I10": "Nhóm chi phí",
}
for cell, value in headers.items():
    ws[cell] = value
    ws[cell].fill = header_fill
    ws[cell].font = header_font
    ws[cell].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws[cell].border = border

analysis_rows = [12, 13, 14, 15, 16]
development_rows = [19, 21, 22, 23, 24, 26, 27, 28, 29, 30, 31]
qa_rows = [34, 35, 36]
uat_rows = [37]

rates = {}
for row in analysis_rows:
    rates[row] = (2_200_000, "Phân tích/thiết kế")
for row in development_rows:
    rates[row] = (2_500_000, "Phát triển")
for row in qa_rows:
    rates[row] = (1_800_000, "Kiểm thử/tối ưu")
for row in uat_rows:
    rates[row] = (2_000_000, "Nghiệm thu/bàn giao")

for row, (rate, group) in rates.items():
    ws.cell(row, 7).value = rate
    ws.cell(row, 8).value = f'=IFERROR(E{row}*G{row},0)'
    ws.cell(row, 9).value = group
    ws.cell(row, 7).number_format = '#,##0" đ"'
    ws.cell(row, 8).number_format = '#,##0" đ"'
    for col in range(7, 10):
        cell = ws.cell(row, col)
        cell.border = border
        cell.alignment = Alignment(horizontal="right" if col in (7, 8) else "left", vertical="center", wrap_text=True)

summary_rows = {
    17: ("Tổng phát triển", development_rows),
    32: ("Tổng trước dự phòng", analysis_rows + development_rows + qa_rows + uat_rows),
    48: ("Tổng manday", analysis_rows + development_rows + qa_rows + uat_rows),
    52: ("Dự phòng 10%", []),
    53: ("TỔNG KINH PHÍ ƯỚC LƯỢNG", []),
}

ws["G17"] = "Tổng phát triển"
ws["H17"] = "=SUM(H19:H31)"
ws["I17"] = "Subtotal"

ws["G32"] = "Tạm tính"
ws["H32"] = "=SUM(H12:H16,H17,H34:H37)"
ws["I32"] = "Chưa gồm dự phòng"

ws["G48"] = "Tổng effort"
ws["H48"] = "=SUM(E12:E16,E17,E34:E37)"
ws["I48"] = "manday"

ws["G52"] = "Dự phòng"
ws["H52"] = "=H32*10%"
ws["I52"] = "10%"

ws["G53"] = "Tổng kinh phí"
ws["H53"] = "=H32+H52"
ws["I53"] = "VND"

for row in [17, 32, 48, 52, 53]:
    for col in range(7, 10):
        cell = ws.cell(row, col)
        cell.fill = subtotal_fill if row != 53 else money_fill
        cell.font = Font(bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="right" if col == 8 else "left", vertical="center", wrap_text=True)
    ws.cell(row, 8).number_format = '#,##0" đ"' if row != 48 else '#,##0'

ws["J10"] = "Ghi chú đơn giá"
ws["J11"] = "Đơn giá là ước lượng nội bộ theo manday, có thể thay đổi theo seniority và phạm vi nghiệm thu."
ws["J10"].fill = header_fill
ws["J10"].font = header_font
ws["J10"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws["J11"].alignment = Alignment(wrap_text=True, vertical="top")
ws["J11"].fill = PatternFill("solid", fgColor="FFF2CC")

widths = {
    "G": 20,
    "H": 24,
    "I": 20,
    "J": 42,
}
for column, width in widths.items():
    ws.column_dimensions[column].width = width

for row in range(10, 54):
    ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 18, 22)

expected_subtotal = sum((ws.cell(row, 5).value or 0) * rate for row, (rate, _) in rates.items())
expected_total = expected_subtotal * 1.1
ws["J13"] = f"Tạm tính hiện tại: {expected_subtotal:,.0f} VND"
ws["J14"] = f"Tổng gồm dự phòng 10%: {expected_total:,.0f} VND"
ws["J13"].number_format = "@"
ws["J14"].number_format = "@"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(path)

print(path.resolve())
print(f"subtotal={expected_subtotal:,.0f}")
print(f"total_with_contingency={expected_total:,.0f}")
