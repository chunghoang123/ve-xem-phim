from pathlib import Path
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

path = Path("../completed/Template-Technical-Completed.xlsx")
wb = load_workbook(path)
ws = wb["2. Price"]

analysis_rows = [12, 13, 14, 15, 16]
development_rows = [19, 21, 22, 23, 24, 26, 27, 28, 29, 30, 31]
qa_rows = [34, 35, 36]
uat_rows = [37]

rates = {}
for row in analysis_rows:
    rates[row] = (180_000, "Phân tích/thiết kế")
for row in development_rows:
    rates[row] = (200_000, "Phát triển")
for row in qa_rows:
    rates[row] = (160_000, "Kiểm thử/tối ưu")
for row in uat_rows:
    rates[row] = (180_000, "Nghiệm thu/bàn giao")

for row, (rate, group) in rates.items():
    ws.cell(row, 7).value = rate
    ws.cell(row, 8).value = f'=IFERROR(E{row}*G{row},0)'
    ws.cell(row, 9).value = group
    ws.cell(row, 7).number_format = '#,##0" đ"'
    ws.cell(row, 8).number_format = '#,##0" đ"'

ws["H17"] = "=SUM(H19:H31)"
ws["H32"] = "=SUM(H12:H16,H17,H34:H37)"
ws["H48"] = "=SUM(E12:E16,E17,E34:E37)"
ws["H52"] = "=H32*10%"
ws["H53"] = "=H32+H52"

subtotal = sum((ws.cell(row, 5).value or 0) * rate for row, (rate, _) in rates.items())
total = subtotal * 1.1

ws["J11"] = "Đơn giá là ước lượng nội bộ ở mức đồ án/sản phẩm MVP, không phải báo giá doanh nghiệp."
ws["J13"] = f"Tạm tính hiện tại: {subtotal:,.0f} VND"
ws["J14"] = f"Tổng gồm dự phòng 10%: {total:,.0f} VND"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(path)

print(path.resolve())
print(f"subtotal={subtotal:,.0f}")
print(f"total_with_contingency={total:,.0f}")
