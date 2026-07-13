from pathlib import Path
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

path = Path("Template-Technical.xlsx")
wb = load_workbook(path, data_only=False)

for ws in wb.worksheets:
    print(f"SHEET\t{ws.title}\t{ws.max_row}x{ws.max_column}")
    print("MERGES\t" + ";".join(str(rng) for rng in ws.merged_cells.ranges))
    non_empty = []
    formulas = []
    blanks_with_style = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                non_empty.append(cell.coordinate)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((cell.coordinate, cell.value))
            elif cell.has_style:
                blanks_with_style.append(cell.coordinate)
    print(f"NON_EMPTY_COUNT\t{len(non_empty)}")
    print(f"FORMULA_COUNT\t{len(formulas)}")
    print("FORMULAS_SAMPLE")
    for coord, formula in formulas[:30]:
        print(f"{coord}\t{formula}")
    print("BLANK_STYLED_SAMPLE")
    for coord in blanks_with_style[:80]:
        print(coord)
    print("VALUES_PREVIEW")
    max_rows = min(ws.max_row, 35)
    max_cols = min(ws.max_column, 18)
    for row in ws.iter_rows(min_row=1, max_row=max_rows, min_col=1, max_col=max_cols):
        values = []
        for cell in row:
            value = cell.value
            if value is None:
                values.append("")
            else:
                text = str(value).replace("\n", " ")
                values.append(text[:80])
        if any(values):
            print("\t".join(values))
    print("END_SHEET")
