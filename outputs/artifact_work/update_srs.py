from pathlib import Path
from docx import Document


SRC = Path("outputs/artifact_work/SRS_DatVePhimAI_v2_HoanThien_CoSoDo.docx")
OUT = Path("outputs/completed/SRS_DatVePhimAI_v2.1_CapNhatTheoHeThong.docx")


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_exact(doc, old, new):
    for p in doc.paragraphs:
        if p.text.strip() == old:
            set_text(p, new)


def replace_contains(doc, needle, new):
    for p in doc.paragraphs:
        if needle in p.text:
            set_text(p, new)


def table_for_uc(doc, uc):
    for table in doc.tables:
        if len(table.rows) and len(table.rows[0].cells) >= 2:
            if table.cell(0, 0).text.strip() == "Use Case ID" and table.cell(0, 1).text.strip() == uc:
                return table
    return None


def set_table_field(table, label, value):
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            set_text(row.cells[1].paragraphs[0], value)
            return


doc = Document(SRC)

replace_exact(doc, "(Smart Cinema Booking with AI Chatbot – Function Calling)", "(Smart Cinema Booking with CineX Conversational Assistant)")
replace_exact(doc, "Ngày ban hành : 08/07/2026", "Ngày cập nhật : 13/07/2026")
replace_exact(doc, "Phiên bản : 2.0", "Phiên bản : 2.1")
replace_contains(doc, "Phạm vi triển khai giới hạn", "Tài liệu bao gồm ERD, use case, luồng xử lý, trạng thái, phân quyền, site map và đặc tả 17 use case. Phiên bản 2.1 đối chiếu với mã nguồn CineX hiện tại: hỗ trợ nhiều rạp, giữ ghế có thời hạn, đặt vé, bắp nước, thanh toán PayOS, vé QR, quản trị và trợ lý hội thoại; các tích hợp TMDB/Google Places phụ thuộc cấu hình dịch vụ ngoài.")
replace_contains(doc, "Website cho phép người dùng xem thông tin phim", "Website cho phép người dùng xem phim và lịch chiếu, chọn ghế, giữ ghế, thêm bắp nước, tạo booking và thanh toán qua PayOS. CineX AI hướng dẫn quy trình đặt vé bằng hội thoại có trạng thái; khi thiếu dữ liệu, trợ lý hỏi lại và sử dụng các bộ xử lý nội bộ để truy vấn phim, rạp, suất chiếu, ghế và trạng thái thanh toán.")
replace_contains(doc, "Luồng xử lý cốt lõi “Đặt vé qua AI Chatbot”", "Luồng cốt lõi: người dùng nhập yêu cầu → CineX AI xác định bước hiện tại → truy vấn dữ liệu phim/rạp/suất/ghế → ghi trạng thái hội thoại → chuyển người dùng sang luồng giữ ghế, tạo booking và thanh toán PayOS. Giao diện đặt vé thông thường luôn là luồng dự phòng.")
replace_contains(doc, "Thời gian phản hồi API thông thường", "Mục tiêu phản hồi API nội bộ dưới 2 giây trong điều kiện tải thông thường; cập nhật trạng thái thanh toán dùng polling/SSE. Tích hợp PayOS, TMDB và Google Places có timeout, xử lý lỗi và thông báo dự phòng.")

for p in doc.paragraphs:
    if p.text.strip().startswith("[Chèn hình"):
        set_text(p, "")

replace_exact(doc, "3.5. UC5 — Đổi / Hủy vé qua Chatbot", "3.5. UC5 — Hủy đơn chưa thanh toán qua CineX AI (đổi suất: mở rộng)")
replace_exact(doc, "3.17. UC17 — Cấu hình / giám sát AI Chatbot", "3.17. UC17 — Giám sát hội thoại CineX AI (phạm vi mở rộng)")
replace_exact(doc, "PHẦN 6: LINK ISSUE", "PHẦN 6: ĐỐI CHIẾU TRIỂN KHAI / LINK ISSUE")

uc3 = table_for_uc(doc, "UC3")
set_table_field(uc3, "Mô tả", "Người dùng đặt vé qua trợ lý CineX AI theo hội thoại nhiều bước; hệ thống lưu trạng thái và truy vấn dữ liệu thật của phim, rạp, suất chiếu và ghế.")
set_table_field(uc3, "Luồng cơ bản (Basic Flow)", "1. Người dùng nêu nhu cầu.\n2. CineX AI xác định phim, khu vực/rạp, ngày và suất.\n3. Hệ thống hiển thị ghế khả dụng.\n4. Người dùng chọn ghế; giao diện tạo seat hold.\n5. Tạo booking PENDING_PAYMENT và chuyển sang PayOS.\n6. Webhook xác nhận, phát hành vé QR.")
set_table_field(uc3, "Yêu cầu phi chức năng (Non-Functional Requirement)", "Trạng thái hội thoại được lưu; không tin giá do client gửi; luôn có đường dẫn sang giao diện đặt vé thường khi trợ lý hoặc dịch vụ ngoài lỗi.")

uc4 = table_for_uc(doc, "UC4")
set_table_field(uc4, "Luồng cơ bản (Basic Flow)", "1. Chọn phim, rạp, ngày, suất chiếu.\n2. Chọn ghế; hệ thống kiểm tra xung đột và giữ ghế.\n3. Chọn bắp nước (tùy chọn); giá và tồn kho được kiểm tra ở server.\n4. Tạo booking và bill trong transaction.\n5. Tạo thanh toán PayOS và theo dõi trạng thái.\n6. Khi webhook hợp lệ: xác nhận booking và sinh vé QR.")

uc5 = table_for_uc(doc, "UC5")
set_table_field(uc5, "Mô tả", "Người dùng hủy booking PENDING_PAYMENT do mình sở hữu; đổi suất hoặc hủy vé đã thanh toán là phạm vi mở rộng chưa tự động hóa.")
set_table_field(uc5, "Điều kiện cần (Pre-Condition)", "Đã đăng nhập; booking thuộc người dùng và đang PENDING_PAYMENT.")
set_table_field(uc5, "Điều kiện sau (Post-Condition(s))", "Booking/Bill chuyển CANCELLED; ghế khóa và seat hold được giải phóng.")
set_table_field(uc5, "Ràng buộc (Business Rules)", "Chỉ hủy tự phục vụ trước thanh toán. Vé CONFIRMED cần quy trình hỗ trợ/hoàn tiền riêng ở phiên bản sau.")

uc7 = table_for_uc(doc, "UC7")
set_table_field(uc7, "Mô tả", "Người dùng thanh toán booking PENDING_PAYMENT qua PayOS bằng checkout URL hoặc QR ngân hàng.")
set_table_field(uc7, "Điều kiện cần (Pre-Condition)", "Booking thuộc người dùng, đang PENDING_PAYMENT, còn hạn 15 phút và số tiền khớp bill do server tính.")
set_table_field(uc7, "Điều kiện sau (Post-Condition(s))", "Payment/Bill = PAID; Booking = CONFIRMED; ghế = BOOKED; ticket và QR được phát hành.")
set_table_field(uc7, "Luồng cơ bản (Basic Flow)", "1. Server tạo payment PayOS từ booking.\n2. Người dùng mở checkout/QR.\n3. Client theo dõi trạng thái bằng polling hoặc SSE.\n4. Webhook PayOS được xác minh chữ ký, đối chiếu amount và chống xử lý lặp.\n5. Transaction xác nhận booking, ghế, bill và tạo ticket.")
set_table_field(uc7, "Ràng buộc (Business Rules)", "Không nhận giá từ client; chỉ PayOS webhook đã xác minh mới xác nhận thanh toán; xử lý webhook phải idempotent.")
set_table_field(uc7, "Yêu cầu phi chức năng (Non-Functional Requirement)", "Không lưu dữ liệu thẻ; ghi payment/webhook log; timeout và trạng thái CANCELLED/EXPIRED/FAILED được xử lý rõ ràng.")

uc17 = table_for_uc(doc, "UC17")
set_table_field(uc17, "Mô tả", "Admin xem trạng thái vận hành hội thoại và cấu hình trợ lý; màn hình chuyên dụng là phạm vi mở rộng. Hiện tại hệ thống đã lưu conversation/state và cung cấp API xem/reset hội thoại có xác thực.")
set_table_field(uc17, "Điều kiện sau (Post-Condition(s))", "Không thay đổi dữ liệu nghiệp vụ trong bản hiện tại; cung cấp dữ liệu nền để xây dựng màn hình giám sát ở phiên bản sau.")

# Document control.
meta = doc.tables[0]
for label, value in {"Trạng thái tài liệu": "FINAL – UPDATED", "Phiên bản": "2.1", "Ngày phát hành": "13/07/2026"}.items():
    set_table_field(meta, label, value)
history = doc.tables[1]
if len(history.rows) > 2:
    vals = ["13/07/2026", "Hoàng Văn Chung", "Cập nhật theo mã nguồn: PayOS, seat hold, JWT refresh, concessions, trạng thái triển khai", ""]
    for cell, value in zip(history.rows[2].cells, vals):
        set_text(cell.paragraphs[0], value)

# Implementation status table.
for table in doc.tables:
    if table.cell(0, 0).text.strip() == "Mã UC" and "Issue" in table.cell(0, 1).text:
        statuses = {
            "UC1": ("src/components/AuthProvider.jsx; server/controllers/auth.controller.cjs", "Hoàn thành"),
            "UC3": ("src/components/ChatBox.jsx; server/services/chat.service.cjs", "Hoàn thành mức hội thoại trạng thái"),
            "UC4": ("src/pages/SeatSelection.jsx; src/pages/Checkout.jsx", "Hoàn thành"),
            "UC7": ("server/services/payment.service.cjs; server/services/webhook.service.cjs", "Hoàn thành PayOS"),
            "UC8": ("server/routes/ticket.routes.cjs", "Backend hoàn thành; UI quét đang hoàn thiện"),
            "UC12-16": ("src/pages/admin; server/routes/admin.routes.cjs", "CRUD chính hoàn thành; báo cáo nâng cao đang hoàn thiện"),
        }
        for row in table.rows[1:]:
            key = row.cells[0].text.strip()
            if key in statuses:
                set_text(row.cells[2].paragraphs[0], statuses[key][0])
                set_text(row.cells[3].paragraphs[0], statuses[key][1])

# Add an implementation delta before Part 6 without changing the established style system.
part6 = next(p for p in doc.paragraphs if p.text.strip().startswith("PHẦN 6:"))
items = [
    ("5.6 Đối chiếu với hiện trạng mã nguồn (13/07/2026)", "Heading 2"),
    ("Đã triển khai: PayOS checkout/QR; webhook xác minh chữ ký và idempotency; giữ ghế/booking 15 phút; vé QR; JWT access 15 phút, refresh 7 ngày và blacklist; bắp nước có tồn kho; giá được tính lại ở server.", "List Bullet"),
    ("Đã triển khai: trang người dùng và admin cho phim, rạp, suất chiếu, booking, khuyến mãi, người dùng và bắp nước; CineX AI dùng hội thoại trạng thái và có luồng dự phòng sang giao diện thường.", "List Bullet"),
    ("Phụ thuộc cấu hình ngoài: TMDB, Google Places, Firebase, email/SMS và PayOS. Khi DNS, billing hoặc khóa API không khả dụng, hệ thống dùng dữ liệu local hoặc thông báo lỗi có kiểm soát.", "List Bullet"),
    ("Kiểm thử tự động hiện có: 3/3 bài kiểm thử tính giá bắp nước, chống giả mạo giá client và giới hạn sản phẩm đã đạt ngày 13/07/2026; cần tiếp tục bổ sung kiểm thử tích hợp cơ sở dữ liệu và webhook.", "List Bullet"),
]
for text, style in items:
    p = part6.insert_paragraph_before(text)
    try:
        p.style = style
    except KeyError:
        p.style = "Normal"

doc.core_properties.title = "SRS – Hệ thống đặt vé xem phim thông minh tích hợp CineX AI"
doc.core_properties.author = "Hoàng Văn Chung"
doc.core_properties.comments = "Phiên bản 2.1 cập nhật theo hiện trạng mã nguồn ngày 13/07/2026"
settings = doc.settings._element
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
if settings.find(qn("w:updateFields")) is None:
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
