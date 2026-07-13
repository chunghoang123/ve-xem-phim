import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document


def text_of_paragraph(paragraph):
    return "".join(run.text for run in paragraph.runs)


def paragraph_record(paragraph, index, part):
    fmt = paragraph.paragraph_format
    return {
        "part": part,
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": text_of_paragraph(paragraph),
        "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
    }


def inspect(path):
    doc = Document(path)
    sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    records = []
    for idx, p in enumerate(doc.paragraphs):
        records.append(paragraph_record(p, idx, "document"))

    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for ri, row in enumerate(table.rows):
            cells = []
            for ci, cell in enumerate(row.cells):
                cells.append({
                    "cell": f"r{ri}c{ci}",
                    "text": "\n".join(p.text for p in cell.paragraphs),
                    "styles": [p.style.name if p.style else None for p in cell.paragraphs],
                })
            rows.append(cells)
        tables.append({"index": ti, "rows": rows})

    headers = []
    footers = []
    sections = []
    for si, section in enumerate(doc.sections):
        sections.append({
            "index": si,
            "width": section.page_width,
            "height": section.page_height,
            "orientation": str(section.orientation),
            "top_margin": section.top_margin,
            "bottom_margin": section.bottom_margin,
            "left_margin": section.left_margin,
            "right_margin": section.right_margin,
            "header_distance": section.header_distance,
            "footer_distance": section.footer_distance,
            "start_type": str(section.start_type),
        })
        headers.extend(paragraph_record(p, i, f"header-{si}") for i, p in enumerate(section.header.paragraphs))
        footers.extend(paragraph_record(p, i, f"footer-{si}") for i, p in enumerate(section.footer.paragraphs))

    with zipfile.ZipFile(path) as zf:
        package = [
            {
                "name": info.filename,
                "size": info.file_size,
                "sha256": hashlib.sha256(zf.read(info.filename)).hexdigest(),
            }
            for info in zf.infolist()
        ]

    return {
        "path": str(path.resolve()),
        "sha256": sha256,
        "paragraphs": records,
        "tables": tables,
        "headers": headers,
        "footers": footers,
        "sections": sections,
        "inline_shapes": len(doc.inline_shapes),
        "package": package,
        "core_properties": {
            "title": doc.core_properties.title,
            "subject": doc.core_properties.subject,
            "author": doc.core_properties.author,
            "keywords": doc.core_properties.keywords,
            "comments": doc.core_properties.comments,
        },
    }


for filename in sys.argv[1:]:
    path = Path(filename)
    data = inspect(path)
    out = path.with_suffix(".inventory.json")
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out)
